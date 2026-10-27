"""Generate frameworks/iso-27001-2022/mappings.json from the official OLIR crosswalk.

Source: NIST OLIR informative reference #155, "800-53-Rev5-to-ISO 27001-2022"
(Final, 2023-11-13):
    https://csrc.nist.gov/csrc/media/Projects/olir/documents/submissions/sp800-53r5-to-iso-27001-mapping-2022-OLIR-2023-10-12-UPDATED.xlsx

The workbook is focal=800-53 → reference=ISO; we invert it to ship edges
toward the hub (ISO control → 800-53). Annex A references carry the `A.`
prefix and match our catalog; bare numbers are ISO management clauses (4-10)
that have no catalog entry and are skipped. NIST's own caveat ("organizations
should not assume ... equivalency") → every edge is `intersects`.

Usage:
    python scripts/build_mappings/build_iso_mappings.py /path/to/mapping.xlsx

Requires: openpyxl (dev-only).
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "frameworks" / "iso-27001-2022" / "mappings.json"

CTRL = re.compile(r"^([A-Z]{2})-0*(\d+)$")  # AC-01 → AC-1


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    import openpyxl

    iso_ids = {
        c["id"]
        for c in json.loads((REPO / "frameworks" / "iso-27001-2022" / "catalog.json").read_text())["controls"]
    }
    n53_ids = {
        c["id"]
        for c in json.loads((REPO / "frameworks" / "nist-800-53-r5" / "catalog.json").read_text())["controls"]
    }

    wb = openpyxl.load_workbook(sys.argv[1])
    mappings = []
    seen = set()
    skipped_clauses = 0
    for sheet in wb.sheetnames:
        if sheet == "Definitions":
            continue
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            focal = str(row[0] or "").strip()  # 800-53, zero-padded
            ref = str(row[3] or "").strip()    # ISO element
            m = CTRL.match(focal)
            if not m or not ref:
                continue
            target = f"{m.group(1)}-{m.group(2)}"
            if target not in n53_ids:
                continue
            if not ref.startswith("A."):
                skipped_clauses += 1
                continue
            if ref not in iso_ids:
                continue
            key = (ref, target)
            if key in seen:
                continue
            seen.add(key)
            mappings.append({
                "control_id": ref,
                "target_framework": "nist-800-53-r5",
                "target_control": target,
                "relationship": "intersects",
                "confidence": "official",
                "rationale": f"SP 800-53r5 ↔ ISO/IEC 27001:2022 crosswalk (OLIR #155): {ref} ↔ {target}",
            })

    OUT.write_text(json.dumps({
        "framework_id": "iso-27001-2022",
        "version": "1.0",
        "source": "NIST OLIR #155 — SP 800-53 Rev 5 to ISO/IEC 27001:2022 crosswalk (Final 2023-11-13), inverted toward the hub",
        "mappings": mappings,
    }, indent=2) + "\n")
    covered = len({m["control_id"] for m in mappings})
    print(f"wrote {OUT.relative_to(REPO)}: {len(mappings)} mappings covering {covered} ISO Annex A controls")
    print(f"skipped {skipped_clauses} management-clause references (no catalog entry)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
